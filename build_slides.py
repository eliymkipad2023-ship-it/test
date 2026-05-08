"""
Vol.16: AIデザイン作成 ― GPT Image 2でデザインを自動化する
Keynote編集可能な.pptxスライドを生成する。
デザイン: 白背景 × 紺色アクセント、シンプル構成。
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from copy import deepcopy
from lxml import etree

# ===== カラーパレット =====
NAVY        = RGBColor(0x14, 0x2A, 0x5C)   # メイン紺
NAVY_DARK   = RGBColor(0x0B, 0x1B, 0x3F)   # 強アクセント
NAVY_LIGHT  = RGBColor(0x3D, 0x5A, 0x80)   # 補助
NAVY_PALE   = RGBColor(0xE8, 0xEE, 0xF7)   # ごく薄い紺（行ストライプ）
GRAY_TEXT   = RGBColor(0x33, 0x33, 0x33)
GRAY_LINE   = RGBColor(0xC8, 0xCF, 0xDC)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_GOLD = RGBColor(0xC8, 0xA8, 0x4B)   # ⚠️マーク等のワンポイント

# ===== スライドサイズ（16:9）=====
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

BLANK = prs.slide_layouts[6]  # 完全空白レイアウト

# ============================================================
# 共通ユーティリティ
# ============================================================
def add_slide():
    return prs.slides.add_slide(BLANK)

def set_fill(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb

def set_no_line(shape):
    shape.line.fill.background()

def set_line(shape, rgb, width_pt=1.0):
    shape.line.color.rgb = rgb
    shape.line.width = Pt(width_pt)

def add_rect(slide, left, top, width, height, fill=WHITE, line=None, line_w=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    set_fill(s, fill)
    if line is None:
        set_no_line(s)
    else:
        set_line(s, line, line_w)
    s.shadow.inherit = False
    return s

def add_round_rect(slide, left, top, width, height, fill=WHITE, line=None, line_w=1.0, radius=0.05):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    s.adjustments[0] = radius
    set_fill(s, fill)
    if line is None:
        set_no_line(s)
    else:
        set_line(s, line, line_w)
    s.shadow.inherit = False
    return s

def add_text(slide, left, top, width, height, text,
             size=18, bold=False, color=GRAY_TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font='Hiragino Sans'):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = text.split('\n') if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return tb

def add_runs(slide, left, top, width, height, runs,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font='Hiragino Sans'):
    """runs: list of (text, size, bold, color)"""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    for i, (text, size, bold, color) in enumerate(runs):
        if text == '__NEWLINE__':
            p = tf.add_paragraph()
            p.alignment = align
            continue
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb

# ===== 共通装飾 =====
def add_header(slide, vol="Vol.16", section_label=None, page_num=None, total=40):
    # 上部ヘッダー帯（細い）
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.32), fill=NAVY)
    # 左上：Volラベル
    add_text(slide, Inches(0.35), Inches(0.04), Inches(4.5), Inches(0.28),
             f"{vol} ｜ AIデザイン作成", size=11, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    # 右上：セクション
    if section_label:
        add_text(slide, Inches(7.0), Inches(0.04), Inches(5.0), Inches(0.28),
                 section_label, size=11, bold=False, color=WHITE,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # ページ番号（右下）
    if page_num is not None:
        add_text(slide, Inches(12.3), Inches(7.10), Inches(0.95), Inches(0.3),
                 f"{page_num} / {total}", size=10, color=NAVY_LIGHT,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

def add_title_block(slide, title, subtitle=None):
    # タイトル本体（紺色、太字、左ボーダー）
    add_rect(slide, Inches(0.5), Inches(0.65), Inches(0.18), Inches(0.85), fill=NAVY)
    add_text(slide, Inches(0.78), Inches(0.55), Inches(11.5), Inches(0.7),
             title, size=30, bold=True, color=NAVY_DARK,
             anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(0.78), Inches(1.20), Inches(11.5), Inches(0.4),
                 subtitle, size=13, bold=False, color=NAVY_LIGHT,
                 anchor=MSO_ANCHOR.TOP)
    # 区切り線
    add_rect(slide, Inches(0.5), Inches(1.62), Inches(12.3), Inches(0.02), fill=NAVY_LIGHT)

def message_box(slide, top, text, left=Inches(0.5), width=Inches(12.3), height=Inches(0.65)):
    """強調メッセージ枠（薄紺背景＋紺左バー）"""
    add_rect(slide, left, top, Inches(0.12), height, fill=NAVY)
    add_rect(slide, left + Inches(0.12), top, width - Inches(0.12), height, fill=NAVY_PALE)
    add_text(slide, left + Inches(0.32), top, width - Inches(0.5), height,
             text, size=14, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)

# ===== 表（テーブル）作成ヘルパー =====
def add_table(slide, left, top, width, height, header, rows,
              col_widths=None, header_h=Inches(0.5), row_h=None,
              header_size=14, body_size=12, align_cols=None):
    n_cols = len(header)
    n_rows = len(rows) + 1
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = tbl_shape.table

    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            table.columns[i].width = Emu(int(width * w / total))

    table.rows[0].height = header_h
    if row_h:
        for r in range(1, n_rows):
            table.rows[r].height = row_h

    # ヘッダー
    for c, h in enumerate(header):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = h
        r.font.size = Pt(header_size)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = 'Hiragino Sans'

    # ボディ
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY_PALE if ri % 2 == 1 else WHITE
            cell.margin_left = Inches(0.10)
            cell.margin_right = Inches(0.10)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.clear()
            tf.word_wrap = True
            # 改行対応
            lines = val.split('\n') if isinstance(val, str) else [val]
            for li, line in enumerate(lines):
                p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                if align_cols and align_cols[ci] == 'C':
                    p.alignment = PP_ALIGN.CENTER
                elif align_cols and align_cols[ci] == 'R':
                    p.alignment = PP_ALIGN.RIGHT
                else:
                    p.alignment = PP_ALIGN.LEFT
                r = p.add_run()
                r.text = line
                r.font.size = Pt(body_size)
                r.font.bold = (ci == 0 and (align_cols is None or align_cols[0] != 'C'))
                # 1列目を強調するパターンを多くするので小ボールド処理を共通化
                if ci == 0 and align_cols and align_cols[0] == 'C':
                    r.font.bold = True
                r.font.color.rgb = NAVY_DARK if (ci == 0) else GRAY_TEXT
                r.font.name = 'Hiragino Sans'
    return table

# ===== セクション扉スライド =====
def add_section_slide(part_label, big_title, page_num):
    slide = add_slide()
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY)
    # 大きな白い装飾線
    add_rect(slide, Inches(0.8), Inches(2.7), Inches(0.7), Inches(0.08), fill=WHITE)
    add_text(slide, Inches(1.6), Inches(2.55), Inches(10), Inches(0.5),
             part_label, size=18, bold=True, color=ACCENT_GOLD,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.8), Inches(3.1), Inches(11.7), Inches(2.0),
             big_title, size=54, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.TOP)
    add_text(slide, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.4),
             "Vol.16 AIデザイン作成", size=13, color=NAVY_PALE)
    add_text(slide, Inches(12.3), Inches(7.1), Inches(0.95), Inches(0.3),
             f"{page_num} / 40", size=10, color=NAVY_PALE, align=PP_ALIGN.RIGHT)
    return slide

# ============================================================
# スライド1: 表紙
# ============================================================
def slide_01():
    slide = add_slide()
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE)
    # 左側の紺色装飾帯
    add_rect(slide, 0, 0, Inches(0.6), SLIDE_H, fill=NAVY)
    add_rect(slide, Inches(0.7), 0, Inches(0.08), SLIDE_H, fill=NAVY_LIGHT)

    # 上部ラベル
    add_text(slide, Inches(1.2), Inches(1.0), Inches(11), Inches(0.5),
             "Vol.16  ｜  AIデザイン作成", size=20, bold=True, color=NAVY_LIGHT)
    # 装飾線
    add_rect(slide, Inches(1.2), Inches(1.55), Inches(2.0), Inches(0.04), fill=NAVY)

    # メインタイトル
    add_text(slide, Inches(1.2), Inches(1.9), Inches(11.5), Inches(1.5),
             "AIデザイン作成", size=68, bold=True, color=NAVY_DARK)

    # サブタイトル
    add_text(slide, Inches(1.2), Inches(3.5), Inches(11.5), Inches(0.6),
             "GPT Image 2でデザインを自動化する", size=28, bold=True, color=NAVY)
    add_text(slide, Inches(1.2), Inches(4.1), Inches(11.5), Inches(0.5),
             "― Phase 4：AIによる創造とアウトプット", size=18, color=NAVY_LIGHT)

    # 日付バー
    add_rect(slide, Inches(1.2), Inches(5.5), Inches(7.0), Inches(0.85), fill=NAVY_PALE)
    add_rect(slide, Inches(1.2), Inches(5.5), Inches(0.15), Inches(0.85), fill=NAVY)
    add_text(slide, Inches(1.55), Inches(5.5), Inches(6.7), Inches(0.85),
             "📅  2026/5/9（土）  19:00 〜", size=22, bold=True, color=NAVY_DARK,
             anchor=MSO_ANCHOR.MIDDLE)

    # フッター
    add_text(slide, Inches(1.2), Inches(6.85), Inches(11), Inches(0.4),
             "AI活用セミナーシリーズ ― Phase 4", size=12, color=NAVY_LIGHT)

slide_01()

# ============================================================
# スライド2: 今日のゴール
# ============================================================
def slide_02():
    slide = add_slide()
    add_header(slide, page_num=2)
    add_title_block(slide, "今日のゴール", "セミナー終了時に「できる」ようになっていることリスト")

    goals = [
        ("1", "ChatGPTのImages機能（GPT Image 2）で、文字指示から画像を生成できる"),
        ("2", "スライド表紙・挿絵・アイコン・図解と、用途に応じてアスペクト比を使い分けられる"),
        ("3", "プロンプト工房Project経由で、抽象的なイメージから高品質な画像生成プロンプトを作らせられる"),
        ("4", "同一テーマ・同一テイストで複数枚の画像を揃えられる"),
        ("5", "生成後にチャットで追加指示を出して、色・要素・配置を編集できる"),
        ("6", "生成した画像をiPadの写真アプリ／ファイルアプリに保存できる"),
        ("7", "生成画像をBase64でClaudeのHTMLドキュメントに埋め込める（Vol.14連携）"),
        ("8", "GPT Image 2の苦手領域（日本語文字・透過非対応）を理解し、白背景生成＋被写体抜き出しで補完できる"),
    ]

    top = Inches(1.85)
    row_h = Inches(0.62)
    for i, (num, text) in enumerate(goals):
        y = top + row_h * i
        # 番号サークル
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), y + Inches(0.05), Inches(0.5), Inches(0.5))
        set_fill(circle, NAVY)
        set_no_line(circle)
        tf = circle.text_frame
        tf.margin_left = 0; tf.margin_right = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = num
        r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
        r.font.name = 'Hiragino Sans'
        # テキスト
        add_text(slide, Inches(1.3), y, Inches(11.7), row_h,
                 text, size=15, bold=False, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)
        # 細い下線
        add_rect(slide, Inches(1.3), y + row_h - Inches(0.03), Inches(11.7), Inches(0.01), fill=GRAY_LINE)

slide_02()

# ============================================================
# スライド3: 今日の流れ
# ============================================================
def slide_03():
    slide = add_slide()
    add_header(slide, page_num=3)
    add_title_block(slide, "今日の流れ", "全体マップ ― Part 0 → Part 4の流れで進めます")

    # タイムラインバー
    parts = [
        ("Part 0", "始める前に",   "10分", NAVY_LIGHT),
        ("Part 1", "基本操作",     "35分", NAVY),
        ("Part 2", "プロンプト工房", "30分", NAVY),
        ("Part 3", "複数枚＆編集",  "20分", NAVY),
        ("Part 4", "Claude連携",   "20分", NAVY_DARK),
    ]
    bar_top = Inches(1.95)
    bar_left = Inches(0.6)
    bar_w = Inches(12.1)
    bar_h = Inches(0.65)
    seg_w = bar_w / 5
    for i, (p, label, t, c) in enumerate(parts):
        x = bar_left + seg_w * i
        add_rect(slide, x, bar_top, seg_w - Inches(0.05), bar_h, fill=c)
        # 矢印（最後以外）
        if i < 4:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, x + seg_w - Inches(0.05),
                                           bar_top, Inches(0.18), bar_h)
            set_fill(arrow, c)
            set_no_line(arrow)
            arrow.rotation = 90
        add_text(slide, x, bar_top, seg_w - Inches(0.05), bar_h,
                 f"{p}  ／  {t}", size=14, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # 詳細表
    header = ["パート", "内容", "時間目安"]
    rows = [
        ("Part 0", "GPT Image 2を始める前に",                       "10分"),
        ("Part 1", "ギャラリー → 基本操作（生成・保存・抜き出し）",   "35分"),
        ("Part 2", "プロンプト工房 ― Projectにプロンプトを書かせる", "30分"),
        ("Part 3", "同一テイストの複数枚 & 生成後の編集",             "20分"),
        ("Part 4", "Claudeに埋め込む ― Vol.14連携",                 "20分"),
    ]
    add_table(slide, Inches(0.6), Inches(2.95), Inches(12.1), Inches(2.95),
              header, rows, col_widths=[2, 8, 2],
              header_h=Inches(0.45), row_h=Inches(0.5),
              header_size=14, body_size=14,
              align_cols=['C', 'L', 'C'])

    # 注釈
    message_box(slide, Inches(6.15),
                "※ Part 1冒頭の「ギャラリーツアー」で、今日学ぶ操作で何が作れるかを先に体感してから技術解説に入ります",
                left=Inches(0.6), width=Inches(12.1))

slide_03()

# ============================================================
# スライド4: セクション ― Part 0
# ============================================================
add_section_slide("PART 0", "GPT Image 2を始める前に", 4)

# ============================================================
# スライド5: 利用できるChatGPTプラン×モード対応表
# ============================================================
def slide_05():
    slide = add_slide()
    add_header(slide, section_label="Part 0｜始める前に", page_num=5)
    add_title_block(slide, "利用できるChatGPTプラン × モード対応表",
                    "どのプランでも画像生成は使える ― モードの違いが少しある")

    header = ["プラン", "利用可否", "利用可能モード"]
    rows = [
        ("Free", "◯", "Instantのみ"),
        ("Go",   "◯", "Instantのみ"),
        ("Plus", "◯", "Instant ＋ Thinking"),
        ("Pro",  "◯", "Instant ＋ Thinking"),
    ]
    add_table(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(2.7),
              header, rows, col_widths=[3, 3, 5.3],
              header_h=Inches(0.5), row_h=Inches(0.55),
              header_size=15, body_size=15,
              align_cols=['C', 'C', 'C'])

    # アイコン的補助
    add_round_rect(slide, Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.2),
                   fill=NAVY_PALE, line=NAVY, line_w=1.0, radius=0.15)
    add_runs(slide, Inches(1.3), Inches(5.05), Inches(11.0), Inches(1.1),
             [
                ("メッセージ：", 16, True, NAVY_DARK),
                ("スライド表紙・挿絵レベルなら ", 16, False, GRAY_TEXT),
                ("無料プランのInstantモードで十分", 16, True, NAVY_DARK),
                ("。", 16, False, GRAY_TEXT),
                ("__NEWLINE__", 0, False, GRAY_TEXT),
                ("Thinkingは「文字が多い・複雑なレイアウト」のときに差が出る。", 15, False, GRAY_TEXT),
             ], anchor=MSO_ANCHOR.MIDDLE)

    # 視覚補足：Instant vs Thinking
    add_text(slide, Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.4),
             "▼ モードの使い分け", size=12, bold=True, color=NAVY)
    # Instant
    add_rect(slide, Inches(1.0), Inches(6.75), Inches(0.3), Inches(0.5), fill=NAVY_LIGHT)
    add_text(slide, Inches(1.4), Inches(6.75), Inches(5.5), Inches(0.5),
             "Instant：軽快・速い・多くの用途で十分", size=13, color=GRAY_TEXT,
             anchor=MSO_ANCHOR.MIDDLE)
    # Thinking
    add_rect(slide, Inches(7.0), Inches(6.75), Inches(0.3), Inches(0.5), fill=NAVY_DARK)
    add_text(slide, Inches(7.4), Inches(6.75), Inches(5.5), Inches(0.5),
             "Thinking：文字多め・複雑レイアウト時に有利", size=13, color=GRAY_TEXT,
             anchor=MSO_ANCHOR.MIDDLE)

slide_05()

# ============================================================
# スライド6: GPT Image 2でできること
# ============================================================
def slide_06():
    slide = add_slide()
    add_header(slide, section_label="Part 0｜始める前に", page_num=6)
    add_title_block(slide, "GPT Image 2でできること", "大きく分けて「作る」「直す」の2方向")

    # 2カラム概念図
    # A 作成
    box_top = Inches(2.0)
    box_h = Inches(4.6)
    box_w = Inches(5.9)

    add_round_rect(slide, Inches(0.6), box_top, box_w, box_h,
                   fill=WHITE, line=NAVY, line_w=2.5, radius=0.05)
    add_rect(slide, Inches(0.6), box_top, box_w, Inches(0.7), fill=NAVY)
    add_text(slide, Inches(0.6), box_top, box_w, Inches(0.7),
             "A. 画像を作成する", size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    items_a = [
        ("テキスト（プロンプト）から画像を生成",        "📝 → 🖼️"),
        ("参考画像を渡してスタイルを真似た画像を生成", "🖼️ → 🖼️"),
    ]
    for i, (t, icon) in enumerate(items_a):
        y = box_top + Inches(1.0) + Inches(1.6) * i
        add_round_rect(slide, Inches(0.9), y, Inches(5.3), Inches(1.3),
                       fill=NAVY_PALE, line=NAVY_LIGHT, line_w=1.0, radius=0.1)
        add_text(slide, Inches(1.05), y + Inches(0.1), Inches(5.0), Inches(0.5),
                 icon, size=24, bold=True, color=NAVY_DARK)
        add_text(slide, Inches(1.05), y + Inches(0.65), Inches(5.0), Inches(0.55),
                 t, size=14, color=GRAY_TEXT)

    # B 編集
    add_round_rect(slide, Inches(6.85), box_top, box_w, box_h,
                   fill=WHITE, line=NAVY, line_w=2.5, radius=0.05)
    add_rect(slide, Inches(6.85), box_top, box_w, Inches(0.7), fill=NAVY)
    add_text(slide, Inches(6.85), box_top, box_w, Inches(0.7),
             "B. 画像を編集する", size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    items_b = [
        ("生成済みの画像にチャットで追加指示", "「色を青に変えて」など"),
        ("一部だけ選択して修正",                 "ピンポイント編集"),
        ("既存の画像をアップロードして編集",     "手元の画像も対象に"),
    ]
    for i, (t, sub) in enumerate(items_b):
        y = box_top + Inches(0.9) + Inches(1.15) * i
        add_round_rect(slide, Inches(7.15), y, Inches(5.3), Inches(0.95),
                       fill=NAVY_PALE, line=NAVY_LIGHT, line_w=1.0, radius=0.1)
        add_text(slide, Inches(7.3), y + Inches(0.1), Inches(5.0), Inches(0.4),
                 t, size=14, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(7.3), y + Inches(0.5), Inches(5.0), Inches(0.4),
                 sub, size=12, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

slide_06()

# ============================================================
# スライド7: 進化点 ― 従来モデルとの違い
# ============================================================
def slide_07():
    slide = add_slide()
    add_header(slide, section_label="Part 0｜始める前に", page_num=7)
    add_title_block(slide, "進化点 ― 従来モデルとの違い",
                    "「画像にロゴ・タイトルを入れたい」が現実的になった")

    # ① 文字精度
    add_round_rect(slide, Inches(0.6), Inches(1.95), Inches(12.1), Inches(2.3),
                   fill=WHITE, line=NAVY, line_w=2.0, radius=0.05)
    add_rect(slide, Inches(0.6), Inches(1.95), Inches(0.18), Inches(2.3), fill=NAVY)
    add_text(slide, Inches(0.95), Inches(2.0), Inches(11.5), Inches(0.5),
             "①  文字の描写精度が大幅向上", size=20, bold=True, color=NAVY_DARK,
             anchor=MSO_ANCHOR.MIDDLE)
    # Before / After
    add_round_rect(slide, Inches(0.95), Inches(2.65), Inches(5.6), Inches(1.4),
                   fill=NAVY_PALE, line=GRAY_LINE, line_w=1.0, radius=0.06)
    add_text(slide, Inches(1.05), Inches(2.7), Inches(5.4), Inches(0.4),
             "BEFORE（以前）", size=11, bold=True, color=NAVY_LIGHT)
    add_text(slide, Inches(1.05), Inches(3.05), Inches(5.4), Inches(1.0),
             "「Confrence 2026」のように\n誤字・崩れだらけで使い物にならない",
             size=14, color=GRAY_TEXT)

    add_round_rect(slide, Inches(6.85), Inches(2.65), Inches(5.6), Inches(1.4),
                   fill=NAVY, line=NAVY, line_w=1.0, radius=0.06)
    add_text(slide, Inches(6.95), Inches(2.7), Inches(5.4), Inches(0.4),
             "AFTER（今）", size=11, bold=True, color=NAVY_PALE)
    add_text(slide, Inches(6.95), Inches(3.05), Inches(5.4), Inches(1.0),
             "バナー・ポスターで読める文字が出せる\n（日本語はまだ崩れあり）",
             size=14, color=WHITE)

    # ② 多言語
    add_round_rect(slide, Inches(0.6), Inches(4.45), Inches(12.1), Inches(1.55),
                   fill=WHITE, line=NAVY, line_w=2.0, radius=0.05)
    add_rect(slide, Inches(0.6), Inches(4.45), Inches(0.18), Inches(1.55), fill=NAVY)
    add_text(slide, Inches(0.95), Inches(4.5), Inches(11.5), Inches(0.5),
             "②  多言語対応の改善", size=20, bold=True, color=NAVY_DARK,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.95), Inches(5.05), Inches(11.5), Inches(0.9),
             "日本語・韓国語・中国語などの文字も以前より精度高く出せる（特に短いキャッチコピーで効果大）",
             size=14, color=GRAY_TEXT)

    message_box(slide, Inches(6.25),
                "メッセージ：「画像にロゴ・タイトルを入れたい」が現実的になった",
                left=Inches(0.6), width=Inches(12.1))

slide_07()

# ============================================================
# スライド8: セクション ― Part 1
# ============================================================
add_section_slide("PART 1", "GPT Image 2でこんなのが作れる\n→ 基本操作", 8)

# ============================================================
# スライド9: 30秒でこんな画像が作れる
# ============================================================
def slide_09():
    slide = add_slide()
    add_header(slide, section_label="Part 1｜基本操作", page_num=9)
    add_title_block(slide, "今日のセミナー後、こんな画像が30秒で作れる",
                    "デザインスキルもPhotoshopも、英語のプロンプトを覚える必要も一切不要")

    # 大きなアピールバナー
    add_rect(slide, Inches(0.6), Inches(1.95), Inches(12.1), Inches(1.6), fill=NAVY)
    add_text(slide, Inches(0.6), Inches(2.0), Inches(12.1), Inches(0.6),
             "デザインスキルもPhotoshopも、英語のプロンプトを覚える必要も一切不要",
             size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.6), Inches(2.65), Inches(12.1), Inches(0.85),
             "30秒で「使える画像」が出てくる ― 必要なのは日本語のひと言だけ",
             size=14, color=NAVY_PALE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # 概念図：日本語ひと言 → ChatGPT → 画像
    boxes = [
        ("日本語の指示",      "「セミナーの表紙画像を作って」", NAVY_LIGHT),
        ("ChatGPT",          "GPT Image 2",                   NAVY),
        ("使える画像",        "30秒で完成 → そのまま使える",     NAVY_DARK),
    ]
    box_w = Inches(3.7)
    box_h = Inches(1.8)
    arrow_w = Inches(0.4)
    total_w = box_w * 3 + arrow_w * 2
    start_x = (SLIDE_W - total_w) / 2
    y = Inches(3.9)
    for i, (t, sub, c) in enumerate(boxes):
        x = start_x + (box_w + arrow_w) * i
        add_round_rect(slide, x, y, box_w, box_h, fill=WHITE, line=c, line_w=2.5, radius=0.06)
        add_rect(slide, x, y, box_w, Inches(0.55), fill=c)
        add_text(slide, x, y, box_w, Inches(0.55),
                 t, size=15, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + Inches(0.1), y + Inches(0.65), box_w - Inches(0.2), Inches(1.05),
                 sub, size=14, color=GRAY_TEXT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 2:
            ax = x + box_w
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, y + box_h/2 - Inches(0.2),
                                           arrow_w, Inches(0.4))
            set_fill(arrow, NAVY_LIGHT); set_no_line(arrow)

    # 説明文
    add_text(slide, Inches(0.6), Inches(6.0), Inches(12.1), Inches(0.5),
             "▼  これから見せる12枚は、すべて「今日学ぶ操作とプロンプト工房だけ」で生成したもの",
             size=14, bold=True, color=NAVY_DARK)
    add_text(slide, Inches(0.6), Inches(6.45), Inches(12.1), Inches(0.5),
             "    自分の仕事・家庭・趣味のどこで使えそうかをイメージしながら見てください",
             size=14, color=GRAY_TEXT)

slide_09()

# ============================================================
# スライド10〜13: ギャラリー（同フォーマット）
# ============================================================
def gallery_slide(page, num_label, title, items):
    slide = add_slide()
    add_header(slide, section_label="Part 1｜ギャラリー", page_num=page)
    add_title_block(slide, title, f"ギャラリー{num_label}  ／  この後の操作で全部作れます")

    # 3カードレイアウト
    card_w = Inches(4.0)
    card_h = Inches(4.5)
    card_gap = Inches(0.15)
    total_w = card_w * 3 + card_gap * 2
    start_x = (SLIDE_W - total_w) / 2
    y = Inches(2.0)

    for i, (img_label, use) in enumerate(items):
        x = start_x + (card_w + card_gap) * i
        add_round_rect(slide, x, y, card_w, card_h,
                       fill=WHITE, line=NAVY, line_w=1.5, radius=0.04)
        # 画像エリア（ダミーフレーム）
        img_area_h = Inches(2.5)
        add_rect(slide, x + Inches(0.15), y + Inches(0.15), card_w - Inches(0.3), img_area_h,
                 fill=NAVY_PALE, line=NAVY_LIGHT, line_w=1.0)
        # 画像中央のラベル（プレースホルダー）
        add_text(slide, x + Inches(0.15), y + Inches(0.15), card_w - Inches(0.3), img_area_h,
                 "🖼  画像プレースホルダー",
                 size=14, color=NAVY_LIGHT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        # 画像説明
        add_rect(slide, x + Inches(0.15), y + Inches(2.7), card_w - Inches(0.3), Inches(0.04),
                 fill=NAVY)
        add_text(slide, x + Inches(0.2), y + Inches(2.78), card_w - Inches(0.4), Inches(0.7),
                 img_label, size=13, bold=True, color=NAVY_DARK)

        # 用途ラベル
        add_text(slide, x + Inches(0.2), y + Inches(3.45), card_w - Inches(0.4), Inches(0.3),
                 "▶ 使い道", size=11, bold=True, color=NAVY_LIGHT)
        add_text(slide, x + Inches(0.2), y + Inches(3.75), card_w - Inches(0.4), Inches(0.7),
                 use, size=12, color=GRAY_TEXT)

# 10
gallery_slide(10, "①  仕事で使う画像",
              "ギャラリー①  仕事で使う画像",
              [
                ("提案書／企画書の表紙\n（16:9・モダンフラット）", "クライアント提案・社内会議のスライド表紙"),
                ("オフィス会議シーンの挿絵\n（横長・水彩風）",   "研修資料・社内報・ブログの挿絵"),
                ("業務マニュアル用ピクトグラム3点\n（電話／メール／会議）", "マニュアル・FAQの統一感あるアイコン"),
              ])

# 11
gallery_slide(11, "②  家庭・プライベートで使う画像",
              "ギャラリー②  家庭・プライベートで使う画像",
              [
                ("お子さん向け誕生日招待状\n（ポップ・1:1）", "LINE配信・印刷して手渡し"),
                ("桜と「春のご挨拶」季節カード\n（縦長・水彩風）", "親戚・取引先への挨拶状"),
                ("「秋祭りのお知らせ」町内会ポスター\n（手描き風）", "掲示板に貼る・LINEグループで配信"),
              ])

# 12
gallery_slide(12, "③  SNS・趣味で使う画像",
              "ギャラリー③  SNS・趣味で使う画像",
              [
                ("レシピブログのアイキャッチ\n（写真風・16:9）", "ブログ・noteのトップ画像"),
                ("「私の朝活ルーティン」Instagram投稿\n（1:1）", "SNS発信・お店紹介"),
                ("旅行ブログのヘッダー画像\n（風景写真風・横長）", "ブログ・SNSのプロフィールヘッダー"),
              ])

# 13
gallery_slide(13, "④  学習・教育用素材",
              "ギャラリー④  学習・教育用素材",
              [
                ("子どもの自由研究の表紙\n（「私のSDGs研究」・1:1）", "学校提出物・家族共有"),
                ("英単語暗記カードのイラスト\n（ポップ・1:1）", "お子さんの勉強・大人の資格学習・GoodNotes活用"),
                ("資格学習／社内研修用の図解キャラクター", "教える側の素材・OJT資料・社内勉強会"),
              ])

# ============================================================
# スライド14: 最初の画像を生成する（操作手順）
# ============================================================
def slide_14():
    slide = add_slide()
    add_header(slide, section_label="Part 1｜基本操作", page_num=14)
    add_title_block(slide, "最初の画像を生成する（操作手順）",
                    "5ステップで完了 ― 慣れれば30秒")

    steps = [
        ("1", "ChatGPTアプリ（またはWeb版）を開く", "アプリ／ブラウザどちらでもOK"),
        ("2", "新規チャットを開始",                  "サイドバー左上のアイコンから"),
        ("3", "プロンプト入力欄に画像生成を依頼",     "例：「青空の下に立つ柴犬の写真を作って」"),
        ("4", "数秒〜30秒で生成完了",                "Instantモードなら更に速い"),
        ("5", "チャット欄に画像が表示される",         "タップで拡大 → 保存可能"),
    ]
    y = Inches(1.95)
    step_h = Inches(0.95)
    gap = Inches(0.05)
    for i, (n, t, sub) in enumerate(steps):
        yi = y + (step_h + gap) * i
        # ステップ番号
        add_rect(slide, Inches(0.6), yi, Inches(1.0), step_h, fill=NAVY)
        add_text(slide, Inches(0.6), yi, Inches(1.0), step_h,
                 f"STEP {n}", size=14, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # 内容
        add_rect(slide, Inches(1.6), yi, Inches(11.1), step_h, fill=NAVY_PALE)
        add_text(slide, Inches(1.85), yi, Inches(7.5), step_h,
                 t, size=16, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(9.4), yi, Inches(3.2), step_h,
                 sub, size=12, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)
        # 次への矢印
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,
                                           Inches(0.95), yi + step_h - Inches(0.02),
                                           Inches(0.3), Inches(0.1))
            set_fill(arrow, NAVY_LIGHT); set_no_line(arrow)

slide_14()

# ============================================================
# スライド15: アスペクト比の使い分け
# ============================================================
def slide_15():
    slide = add_slide()
    add_header(slide, section_label="Part 1｜基本操作", page_num=15)
    add_title_block(slide, "アスペクト比の使い分け",
                    "用途で選ぶ ― 迷ったら下のうちのどれか")

    # 3つの形を視覚化
    aspects = [
        ("16:9",  Inches(2.6), Inches(1.46), "YouTubeサムネイル\nスライド表紙",   NAVY),
        ("1:1",   Inches(1.7), Inches(1.7),  "Instagram投稿\nアイコン",          NAVY),
        ("9:16",  Inches(1.1), Inches(1.96), "Instagramストーリー\nTikTok",      NAVY),
    ]
    total_w = sum(a[1] for a in aspects) + Inches(1.6)
    start_x = (SLIDE_W - total_w) / 2
    y_top = Inches(1.95)
    max_h = Inches(1.96)
    cur_x = start_x
    for ratio, w, h, use, c in aspects:
        # 縦中央揃え（最大高さに合わせる）
        offset = (max_h - h) / 2
        add_rect(slide, cur_x, y_top + offset, w, h, fill=NAVY_PALE, line=NAVY, line_w=2.0)
        add_text(slide, cur_x, y_top + offset, w, h,
                 ratio, size=30, bold=True, color=NAVY_DARK,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # 用途
        add_text(slide, cur_x - Inches(0.2), y_top + max_h + Inches(0.1), w + Inches(0.4), Inches(0.85),
                 use, size=12, color=GRAY_TEXT,
                 align=PP_ALIGN.CENTER)
        cur_x += w + Inches(0.8)

    # 詳細表
    header = ["用途", "推奨アスペクト比"]
    rows = [
        ("YouTubeサムネイル／スライド表紙",       "16 : 9"),
        ("Instagram投稿／アイコン",              "1 : 1"),
        ("Instagramストーリー／TikTok",          "9 : 16"),
    ]
    add_table(slide, Inches(2.5), Inches(5.0), Inches(8.3), Inches(1.5),
              header, rows, col_widths=[6, 4],
              header_h=Inches(0.4), row_h=Inches(0.36),
              header_size=13, body_size=13,
              align_cols=['L', 'C'])

    message_box(slide, Inches(6.6),
                "メッセージ：プロンプトに直接入れる、または生成後に「16:9で作り直して」と追加指示する",
                left=Inches(0.6), width=Inches(12.1))

slide_15()

# ============================================================
# スライド16: iPadへの画像保存 ― 2つの保存先
# ============================================================
def slide_16():
    slide = add_slide()
    add_header(slide, section_label="Part 1｜基本操作", page_num=16)
    add_title_block(slide, "iPadへの画像保存 ― 2つの保存先",
                    "目的に合わせて選ぶ")

    # 2カラム比較
    col_w = Inches(5.9)
    col_h = Inches(3.3)
    y = Inches(2.0)
    # 写真App
    add_round_rect(slide, Inches(0.6), y, col_w, col_h, fill=WHITE, line=NAVY, line_w=2.5, radius=0.05)
    add_rect(slide, Inches(0.6), y, col_w, Inches(0.7), fill=NAVY)
    add_text(slide, Inches(0.6), y, col_w, Inches(0.7),
             "📷  写真アプリ", size=22, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets_a = [
        "SNS投稿に最適",
        "Keynoteクイック貼り付け",
        "iPhoneとも自動同期（iCloud写真）",
    ]
    for i, b in enumerate(bullets_a):
        yy = y + Inches(0.95) + Inches(0.5) * i
        # 丸ポチ
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), yy + Inches(0.13), Inches(0.18), Inches(0.18))
        set_fill(dot, NAVY); set_no_line(dot)
        add_text(slide, Inches(1.25), yy, col_w - Inches(0.7), Inches(0.5),
                 b, size=14, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # ファイルApp
    add_round_rect(slide, Inches(6.85), y, col_w, col_h, fill=WHITE, line=NAVY, line_w=2.5, radius=0.05)
    add_rect(slide, Inches(6.85), y, col_w, Inches(0.7), fill=NAVY)
    add_text(slide, Inches(6.85), y, col_w, Inches(0.7),
             "📁  ファイルアプリ（iCloud Drive等）", size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets_b = [
        "他端末との共有・同期",
        "Claudeに添付して連携",
        "フォルダで整理・案件別管理",
    ]
    for i, b in enumerate(bullets_b):
        yy = y + Inches(0.95) + Inches(0.5) * i
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.2), yy + Inches(0.13), Inches(0.18), Inches(0.18))
        set_fill(dot, NAVY); set_no_line(dot)
        add_text(slide, Inches(7.5), yy, col_w - Inches(0.7), Inches(0.5),
                 b, size=14, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # 操作フロー
    add_text(slide, Inches(0.6), Inches(5.5), Inches(12.1), Inches(0.4),
             "▼  操作の流れ", size=13, bold=True, color=NAVY_DARK)
    flow_items = ["画像をタップ拡大", "右上の共有アイコン", "「画像を保存」 or 「ファイルに保存」"]
    box_w_f = Inches(3.7)
    arrow_w_f = Inches(0.4)
    total_f = box_w_f * 3 + arrow_w_f * 2
    start_xf = (SLIDE_W - total_f) / 2
    yf = Inches(6.0)
    for i, t in enumerate(flow_items):
        x = start_xf + (box_w_f + arrow_w_f) * i
        add_round_rect(slide, x, yf, box_w_f, Inches(0.85),
                       fill=NAVY_PALE, line=NAVY, line_w=1.5, radius=0.1)
        add_text(slide, x, yf, box_w_f, Inches(0.85),
                 t, size=14, bold=True, color=NAVY_DARK,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 2:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + box_w_f, yf + Inches(0.25),
                                           arrow_w_f, Inches(0.35))
            set_fill(arrow, NAVY); set_no_line(arrow)

slide_16()

# ============================================================
# スライド17: 透過代替テクニック（フローチャート）
# ============================================================
def slide_17():
    slide = add_slide()
    add_header(slide, section_label="Part 1｜基本操作", page_num=17)
    add_title_block(slide, "⚠️ 透過代替テクニック ― 白背景生成＋被写体抜き出し",
                    "GPT Image 2は透過PNG非対応 → iPadのネイティブ機能で補完できる")

    # フローチャート（縦6ステップ）
    steps = [
        "Step 1：プロンプトに「pure white background」を指定して生成",
        "Step 2：写真アプリに保存",
        "Step 3：写真アプリで画像を開いて被写体を長押し",
        "Step 4：被写体だけが浮き上がる（輪郭が光る）",
        "Step 5：「コピー」or「ステッカーを追加」",
        "Step 6：Keynote／Pages／メモなどに貼り付け  →  透過素材として使える",
    ]
    box_w = Inches(11)
    box_h = Inches(0.65)
    gap = Inches(0.12)
    start_y = Inches(1.95)
    start_x = (SLIDE_W - box_w) / 2

    colors = [NAVY_LIGHT, NAVY_LIGHT, NAVY, NAVY, NAVY_DARK, NAVY_DARK]
    for i, (t, c) in enumerate(zip(steps, colors)):
        y = start_y + (box_h + gap) * i
        add_round_rect(slide, start_x, y, box_w, box_h,
                       fill=NAVY_PALE if i % 2 == 0 else WHITE,
                       line=c, line_w=1.5, radius=0.1)
        # 番号バー
        add_rect(slide, start_x, y, Inches(0.18), box_h, fill=c)
        add_text(slide, start_x + Inches(0.4), y, box_w - Inches(0.5), box_h,
                 t, size=14, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        # 矢印
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,
                                           start_x + box_w / 2 - Inches(0.15),
                                           y + box_h + Inches(0.005),
                                           Inches(0.3), Inches(0.11))
            set_fill(arrow, NAVY_LIGHT); set_no_line(arrow)

slide_17()

# ============================================================
# スライド18: 白背景キーワード
# ============================================================
def slide_18():
    slide = add_slide()
    add_header(slide, section_label="Part 1｜基本操作", page_num=18)
    add_title_block(slide, "プロンプトに入れる「白背景キーワード」",
                    "コピペ用 ― この語を1つ含めるだけで抜きやすさが激変")

    keywords = [
        ("pure white background",       "真っ白な背景",            NAVY),
        ("solid white background",      "均一な白背景",            NAVY),
        ("plain background, no scenery","風景なし、無地背景",       NAVY_DARK),
    ]
    y = Inches(2.0)
    box_h = Inches(1.1)
    gap = Inches(0.2)
    for i, (en, ja, c) in enumerate(keywords):
        yi = y + (box_h + gap) * i
        add_round_rect(slide, Inches(0.6), yi, Inches(12.1), box_h,
                       fill=WHITE, line=c, line_w=2.0, radius=0.05)
        add_rect(slide, Inches(0.6), yi, Inches(0.18), box_h, fill=c)
        # 英語キーワード（コード風背景）
        add_rect(slide, Inches(0.95), yi + Inches(0.18), Inches(6.0), Inches(0.74),
                 fill=NAVY_PALE)
        add_text(slide, Inches(1.05), yi + Inches(0.18), Inches(5.85), Inches(0.74),
                 en, size=20, bold=True, color=NAVY_DARK,
                 anchor=MSO_ANCHOR.MIDDLE,
                 font='Menlo')
        # 日本語訳
        add_text(slide, Inches(7.2), yi, Inches(5.4), box_h,
                 "▶  " + ja, size=18, bold=True, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    message_box(slide, Inches(6.4),
                "メッセージ：単色背景（白・グレー・薄い青）で生成すると、iPad側の被写体抜き出しの精度が大きく上がる",
                left=Inches(0.6), width=Inches(12.1))

slide_18()

# ============================================================
# スライド19: iPad被写体抜き出しの手順
# ============================================================
def slide_19():
    slide = add_slide()
    add_header(slide, section_label="Part 1｜基本操作", page_num=19)
    add_title_block(slide, "iPad被写体抜き出しの手順",
                    "iOS/iPadOSの「Visual Look Up」機能を使う")

    steps = [
        ("1", "写真アプリで画像を開く（全画面表示）", ""),
        ("2", "被写体を指で長押し（2〜3秒）",       "押しっぱなしでOK"),
        ("3", "輪郭が白く光ったら、指を離す",       "認識成功のサイン"),
        ("4", "メニューから選ぶ",                    "「コピー」「ステッカーを追加」「共有」"),
    ]
    y = Inches(1.95)
    sh = Inches(0.85)
    gap = Inches(0.1)
    for i, (n, t, sub) in enumerate(steps):
        yi = y + (sh + gap) * i
        # 番号
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), yi, Inches(0.85), Inches(0.85))
        set_fill(circle, NAVY); set_no_line(circle)
        tf = circle.text_frame
        tf.margin_left = 0; tf.margin_right = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = n
        r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = WHITE
        r.font.name = 'Hiragino Sans'
        # 内容
        add_round_rect(slide, Inches(1.6), yi, Inches(11.1), sh,
                       fill=NAVY_PALE, line=NAVY_LIGHT, line_w=1.0, radius=0.08)
        add_text(slide, Inches(1.85), yi, Inches(7.5), sh,
                 t, size=16, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        if sub:
            add_text(slide, Inches(9.4), yi, Inches(3.2), sh,
                     sub, size=12, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # 2つの選択肢
    add_text(slide, Inches(0.6), Inches(5.85), Inches(12.1), Inches(0.4),
             "▼  使い分け", size=14, bold=True, color=NAVY_DARK)

    add_round_rect(slide, Inches(0.6), Inches(6.3), Inches(5.9), Inches(0.95),
                   fill=WHITE, line=NAVY, line_w=1.5, radius=0.06)
    add_rect(slide, Inches(0.6), Inches(6.3), Inches(0.15), Inches(0.95), fill=NAVY)
    add_text(slide, Inches(0.85), Inches(6.3), Inches(5.65), Inches(0.4),
             "即貼り付けたい", size=13, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.85), Inches(6.7), Inches(5.65), Inches(0.5),
             "「コピー」 → 別アプリで「ペースト」", size=13, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    add_round_rect(slide, Inches(6.8), Inches(6.3), Inches(5.9), Inches(0.95),
                   fill=WHITE, line=NAVY, line_w=1.5, radius=0.06)
    add_rect(slide, Inches(6.8), Inches(6.3), Inches(0.15), Inches(0.95), fill=NAVY)
    add_text(slide, Inches(7.05), Inches(6.3), Inches(5.65), Inches(0.4),
             "何度も使いたい", size=13, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(7.05), Inches(6.7), Inches(5.65), Inches(0.5),
             "「ステッカーを追加」 → 各アプリで使い回せる", size=13, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

slide_19()

# ============================================================
# スライド20: 演習①
# ============================================================
def exercise_slide(page, title_main, title_sub, steps, checks, time_label, section="Part 1｜演習"):
    slide = add_slide()
    add_header(slide, section_label=section, page_num=page)

    # 演習バッジ
    add_rect(slide, Inches(0.5), Inches(0.65), Inches(0.18), Inches(0.85), fill=ACCENT_GOLD)
    add_runs(slide, Inches(0.78), Inches(0.55), Inches(11.5), Inches(0.7),
             [
                (title_main, 30, True, NAVY_DARK),
                ("   ", 30, False, NAVY_DARK),
                (f"⏱ {time_label}", 18, True, ACCENT_GOLD),
             ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.78), Inches(1.20), Inches(11.5), Inches(0.4),
             title_sub, size=14, bold=False, color=NAVY_LIGHT)
    add_rect(slide, Inches(0.5), Inches(1.62), Inches(12.3), Inches(0.02), fill=NAVY_LIGHT)

    # 左：手順
    add_text(slide, Inches(0.6), Inches(1.85), Inches(7.0), Inches(0.45),
             "▼  手順", size=16, bold=True, color=NAVY_DARK)
    for i, s in enumerate(steps):
        yi = Inches(2.35) + Inches(0.7) * i
        # 番号
        add_rect(slide, Inches(0.6), yi, Inches(0.6), Inches(0.6), fill=NAVY)
        add_text(slide, Inches(0.6), yi, Inches(0.6), Inches(0.6),
                 str(i + 1), size=20, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # 本文
        add_rect(slide, Inches(1.25), yi, Inches(6.5), Inches(0.6), fill=NAVY_PALE)
        add_text(slide, Inches(1.4), yi, Inches(6.3), Inches(0.6),
                 s, size=13, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # 右：チェックポイント
    add_text(slide, Inches(8.0), Inches(1.85), Inches(4.7), Inches(0.45),
             "▼  チェックポイント", size=16, bold=True, color=NAVY_DARK)
    add_round_rect(slide, Inches(8.0), Inches(2.35), Inches(4.7), Inches(4.9),
                   fill=NAVY_PALE, line=NAVY, line_w=1.5, radius=0.05)
    for i, c in enumerate(checks):
        yi = Inches(2.55) + Inches(0.6) * i
        # チェックボックス
        cb = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.2), yi + Inches(0.05),
                                    Inches(0.32), Inches(0.32))
        set_fill(cb, WHITE); set_line(cb, NAVY, 1.5)
        add_text(slide, Inches(8.6), yi, Inches(4.0), Inches(0.4),
                 c, size=13, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)

slide20_steps = [
    "自分の好きなテーマで人物 or キャラクターを生成（プロンプトに必ず「pure white background」）",
    "生成画像を写真アプリに保存",
    "写真アプリで画像を開いて被写体を長押し → 抜き出し",
    "Keynoteで新規スライドを作成し、抜き出した被写体をペースト",
]
slide20_checks = [
    "1枚生成できた",
    "写真アプリに保存できた",
    "被写体抜き出しでKeynoteに貼り付けられた",
]
exercise_slide(20, "演習①",
               "生成 → 保存 → 被写体抜き出しの一連の流れを体験する",
               slide20_steps, slide20_checks, "10分")

# ============================================================
# スライド21: セクション ― Part 2
# ============================================================
add_section_slide("PART 2", "プロンプト工房 ― Projectに\nプロンプトを書かせる", 21)

# ============================================================
# スライド22: なぜ自分で書かないのがベストか
# ============================================================
def slide_22():
    slide = add_slide()
    add_header(slide, section_label="Part 2｜プロンプト工房", page_num=22)
    add_title_block(slide, "なぜ「自分で書かない」のがベストか",
                    "良いプロンプトには10要素が必要 → 毎回思い出すのは無理")

    # 10要素を10個のチップで
    elems = [
        "用途", "背景・場面", "被写体", "スタイル", "色味・カラーパレット",
        "構図・カメラアングル", "ライティング", "画面内の文字", "アスペクト比", "制約",
    ]
    add_text(slide, Inches(0.6), Inches(1.85), Inches(12.1), Inches(0.4),
             "▼  良いプロンプトに必要な10要素", size=14, bold=True, color=NAVY_DARK)

    chip_w = Inches(2.3)
    chip_h = Inches(0.55)
    chip_gap_x = Inches(0.1)
    chip_gap_y = Inches(0.12)
    cols = 5
    start_x = Inches(0.6)
    start_y = Inches(2.35)
    for i, e in enumerate(elems):
        r = i // cols
        c = i % cols
        x = start_x + (chip_w + chip_gap_x) * c
        y = start_y + (chip_h + chip_gap_y) * r
        add_round_rect(slide, x, y, chip_w, chip_h,
                       fill=NAVY, line=NAVY, line_w=1.0, radius=0.3)
        add_text(slide, x, y, chip_w, chip_h,
                 e, size=12, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Before / After
    y_ba = Inches(4.05)
    box_w = Inches(5.95)
    box_h = Inches(2.85)
    # Before
    add_round_rect(slide, Inches(0.6), y_ba, box_w, box_h,
                   fill=WHITE, line=NAVY_LIGHT, line_w=1.5, radius=0.05)
    add_rect(slide, Inches(0.6), y_ba, box_w, Inches(0.55), fill=NAVY_LIGHT)
    add_text(slide, Inches(0.6), y_ba, box_w, Inches(0.55),
             "Before：自分で書く", size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets_before = [
        "コツを毎回思い出す",
        "10項目を書き並べる",
        "結果：質が安定しない",
    ]
    for i, b in enumerate(bullets_before):
        yy = y_ba + Inches(0.7) + Inches(0.65) * i
        add_text(slide, Inches(0.95), yy, Inches(5.5), Inches(0.5),
                 f"✕  {b}", size=14, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # After
    add_round_rect(slide, Inches(6.75), y_ba, box_w, box_h,
                   fill=WHITE, line=NAVY, line_w=2.0, radius=0.05)
    add_rect(slide, Inches(6.75), y_ba, box_w, Inches(0.55), fill=NAVY)
    add_text(slide, Inches(6.75), y_ba, box_w, Inches(0.55),
             "After：Projectに書かせる", size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets_after = [
        "「○○作りたい」と話すだけ",
        "Projectが質問してくれる",
        "完成プロンプトをコピーするだけ",
    ]
    for i, b in enumerate(bullets_after):
        yy = y_ba + Inches(0.7) + Inches(0.65) * i
        add_text(slide, Inches(7.1), yy, Inches(5.5), Inches(0.5),
                 f"✓  {b}", size=14, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)

slide_22()

# ============================================================
# スライド23: ChatGPT Projectsの基本（概念図）
# ============================================================
def slide_23():
    slide = add_slide()
    add_header(slide, section_label="Part 2｜プロンプト工房", page_num=23)
    add_title_block(slide, "ChatGPT Projectsの基本",
                    "Projectには3つの設定枠がある")

    # 中央に「Project」の枠 → 3要素を内包する図
    proj_x = Inches(0.6)
    proj_y = Inches(2.0)
    proj_w = Inches(12.1)
    proj_h = Inches(4.4)
    add_round_rect(slide, proj_x, proj_y, proj_w, proj_h,
                   fill=WHITE, line=NAVY, line_w=2.5, radius=0.04)
    # ヘッダー
    add_rect(slide, proj_x, proj_y, proj_w, Inches(0.6), fill=NAVY)
    add_text(slide, proj_x, proj_y, proj_w, Inches(0.6),
             "📂  Project", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # 3要素
    items = [
        ("1", "システムプロンプト\n（指示）",
         "このProjectでChatGPTがどう振る舞うか"),
        ("2", "知識ベース\n（ファイル）",
         "このProject内のチャットで参照される資料"),
        ("3", "会話履歴",
         "このProject内の会話だけがまとまる"),
    ]
    sub_w = Inches(3.7)
    sub_h = Inches(2.95)
    sub_gap = Inches(0.2)
    sub_total = sub_w * 3 + sub_gap * 2
    sub_start = proj_x + (proj_w - sub_total) / 2
    sub_y = proj_y + Inches(0.85)
    for i, (n, t, sub) in enumerate(items):
        x = sub_start + (sub_w + sub_gap) * i
        add_round_rect(slide, x, sub_y, sub_w, sub_h,
                       fill=NAVY_PALE, line=NAVY_LIGHT, line_w=1.5, radius=0.05)
        # 番号サークル
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + sub_w/2 - Inches(0.3), sub_y + Inches(0.2),
                                        Inches(0.6), Inches(0.6))
        set_fill(circle, NAVY); set_no_line(circle)
        tf = circle.text_frame; tf.margin_left = 0; tf.margin_right = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = n
        r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = WHITE
        r.font.name = 'Hiragino Sans'
        # タイトル
        add_text(slide, x, sub_y + Inches(0.95), sub_w, Inches(0.85),
                 t, size=16, bold=True, color=NAVY_DARK,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # 説明
        add_text(slide, x + Inches(0.15), sub_y + Inches(1.95), sub_w - Inches(0.3), Inches(0.9),
                 sub, size=12, color=GRAY_TEXT,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)

    message_box(slide, Inches(6.65),
                "メッセージ：一度設定すれば、毎回コツを伝え直す必要がない",
                left=Inches(0.6), width=Inches(12.1))

slide_23()

# ============================================================
# スライド24: Project プラン別ファイル数上限
# ============================================================
def slide_24():
    slide = add_slide()
    add_header(slide, section_label="Part 2｜プロンプト工房", page_num=24)
    add_title_block(slide, "Project ― プラン別のファイル数上限",
                    "プロンプト工房は配布物①の1ファイルだけ ＝ 無料プランでも問題なし")

    header = ["プラン", "Project内のファイルアップロード上限"]
    rows = [
        ("Free",                       " 5  ファイル"),
        ("Plus",                       "25  ファイル"),
        ("Pro / Business / Enterprise", "40  ファイル"),
    ]
    add_table(slide, Inches(1.5), Inches(2.05), Inches(10.3), Inches(2.1),
              header, rows, col_widths=[5, 5],
              header_h=Inches(0.5), row_h=Inches(0.5),
              header_size=15, body_size=15,
              align_cols=['C', 'C'])

    # 視覚的に「1ファイルでOK」
    add_round_rect(slide, Inches(2.5), Inches(4.5), Inches(8.3), Inches(1.6),
                   fill=NAVY, line=NAVY, line_w=2.0, radius=0.05)
    add_text(slide, Inches(2.5), Inches(4.55), Inches(8.3), Inches(0.5),
             "🗂  プロンプト工房に必要なファイル数", size=14, bold=True, color=NAVY_PALE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(2.5), Inches(5.05), Inches(8.3), Inches(0.95),
             "1 ファイル（配布物①）", size=42, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    message_box(slide, Inches(6.4),
                "メッセージ：プロンプト工房は配布物①の1ファイルだけで動く ＝ 無料プランでも問題なし",
                left=Inches(0.6), width=Inches(12.1))

slide_24()

# ============================================================
# スライド25: プロンプト工房を作る Step 1〜2
# ============================================================
def slide_25():
    slide = add_slide()
    add_header(slide, section_label="Part 2｜プロンプト工房", page_num=25)
    add_title_block(slide, "プロンプト工房Projectを作る ― Step 1〜2",
                    "①新規Project作成 → ②システムプロンプト設定")

    # Step 1
    add_round_rect(slide, Inches(0.6), Inches(1.95), Inches(12.1), Inches(2.4),
                   fill=WHITE, line=NAVY, line_w=2.0, radius=0.05)
    add_rect(slide, Inches(0.6), Inches(1.95), Inches(2.0), Inches(2.4), fill=NAVY)
    add_text(slide, Inches(0.6), Inches(1.95), Inches(2.0), Inches(0.55),
             "Step 1", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.6), Inches(2.5), Inches(2.0), Inches(1.85),
             "新規Projectを\n作成", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets_1 = [
        "ChatGPTのサイドバーから「Projects」セクションを開く",
        "「+ New Project」（or「+ 新規プロジェクト」）をタップ",
        "Project名を入力（例：「プロンプト工房」）",
    ]
    for i, b in enumerate(bullets_1):
        yy = Inches(2.1) + Inches(0.65) * i
        add_text(slide, Inches(2.85), yy, Inches(9.5), Inches(0.6),
                 f"●  {b}", size=14, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # Step 2
    add_round_rect(slide, Inches(0.6), Inches(4.5), Inches(12.1), Inches(2.5),
                   fill=WHITE, line=NAVY, line_w=2.0, radius=0.05)
    add_rect(slide, Inches(0.6), Inches(4.5), Inches(2.0), Inches(2.5), fill=NAVY)
    add_text(slide, Inches(0.6), Inches(4.5), Inches(2.0), Inches(0.55),
             "Step 2", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.6), Inches(5.05), Inches(2.0), Inches(1.95),
             "システム\nプロンプト\nを設定", size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets_2 = [
        "「Instructions」（or「指示」）欄を開く",
        "配布物②のテキストを全文コピーして貼り付け",
        "保存",
    ]
    for i, b in enumerate(bullets_2):
        yy = Inches(4.7) + Inches(0.7) * i
        # 強調：2番目だけ目立たせる
        if i == 1:
            add_rect(slide, Inches(2.85), yy + Inches(0.05), Inches(9.5), Inches(0.55), fill=NAVY_PALE)
            add_text(slide, Inches(3.0), yy, Inches(9.4), Inches(0.65),
                     f"●  配布物②のテキストを全文コピーして貼り付け",
                     size=14, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        else:
            add_text(slide, Inches(2.85), yy, Inches(9.5), Inches(0.65),
                     f"●  {b}", size=14, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

slide_25()

# ============================================================
# スライド26: プロンプト工房を作る Step 3〜4
# ============================================================
def slide_26():
    slide = add_slide()
    add_header(slide, section_label="Part 2｜プロンプト工房", page_num=26)
    add_title_block(slide, "プロンプト工房Projectを作る ― Step 3〜4",
                    "③知識ベースをアップロード → ④動作確認")

    # Step 3
    add_round_rect(slide, Inches(0.6), Inches(1.95), Inches(12.1), Inches(2.4),
                   fill=WHITE, line=NAVY, line_w=2.0, radius=0.05)
    add_rect(slide, Inches(0.6), Inches(1.95), Inches(2.0), Inches(2.4), fill=NAVY)
    add_text(slide, Inches(0.6), Inches(1.95), Inches(2.0), Inches(0.55),
             "Step 3", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.6), Inches(2.5), Inches(2.0), Inches(1.85),
             "知識ベース\n（配布物①）\nをアップロード", size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets_3 = [
        "「Files」（or「ファイル」）欄を開く",
        "配布物①のPDFをアップロード",
        "アップロード完了を確認",
    ]
    for i, b in enumerate(bullets_3):
        yy = Inches(2.1) + Inches(0.65) * i
        if i == 1:
            add_rect(slide, Inches(2.85), yy + Inches(0.05), Inches(9.5), Inches(0.55), fill=NAVY_PALE)
            add_text(slide, Inches(3.0), yy, Inches(9.4), Inches(0.65),
                     f"●  配布物①のPDFをアップロード",
                     size=14, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        else:
            add_text(slide, Inches(2.85), yy, Inches(9.5), Inches(0.65),
                     f"●  {b}", size=14, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # Step 4
    add_round_rect(slide, Inches(0.6), Inches(4.5), Inches(12.1), Inches(2.5),
                   fill=WHITE, line=NAVY, line_w=2.0, radius=0.05)
    add_rect(slide, Inches(0.6), Inches(4.5), Inches(2.0), Inches(2.5), fill=NAVY)
    add_text(slide, Inches(0.6), Inches(4.5), Inches(2.0), Inches(0.55),
             "Step 4", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.6), Inches(5.05), Inches(2.0), Inches(1.95),
             "動作確認", size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets_4 = [
        "このProject内で新規チャットを開始",
        "「画像を作りたいです」と話しかける",
        "Projectがヒアリング質問を返してくれれば成功",
    ]
    for i, b in enumerate(bullets_4):
        yy = Inches(4.7) + Inches(0.7) * i
        if i == 2:
            add_rect(slide, Inches(2.85), yy + Inches(0.05), Inches(9.5), Inches(0.55), fill=NAVY_PALE)
            add_text(slide, Inches(3.0), yy, Inches(9.4), Inches(0.65),
                     f"●  Projectがヒアリング質問を返してくれれば成功",
                     size=14, bold=True, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        else:
            add_text(slide, Inches(2.85), yy, Inches(9.5), Inches(0.65),
                     f"●  {b}", size=14, color=GRAY_TEXT, anchor=MSO_ANCHOR.MIDDLE)

slide_26()

# ============================================================
# スライド27: 使い方フロー
# ============================================================
def slide_27():
    slide = add_slide()
    add_header(slide, section_label="Part 2｜プロンプト工房", page_num=27)
    add_title_block(slide, "使い方フロー",
                    "工房は『プロンプトを作る場所』。画像生成は別チャットで行う")

    # フロー（縦）
    flows = [
        ("あなた",  "「セミナーの表紙画像を作りたい。\nテーマはAI活用」",                     NAVY_LIGHT),
        ("工房",    "「対象は誰向け？」「写真風 or イラスト風？」\n「色味の希望は？」「文字を入れる？」", NAVY),
        ("あなた",  "（質問に答える）",                                                     NAVY_LIGHT),
        ("工房",    "「以下のプロンプトをImages機能に貼ってください」 → 完成プロンプトをコピー", NAVY),
        ("画像生成", "新規チャット（Project外）に貼り付けて画像生成",                         NAVY_DARK),
    ]
    y0 = Inches(1.92)
    box_h = Inches(0.8)
    box_gap = Inches(0.15)
    for i, (who, text, c) in enumerate(flows):
        y = y0 + (box_h + box_gap) * i
        # ロール
        add_rect(slide, Inches(0.7), y, Inches(2.0), box_h, fill=c)
        add_text(slide, Inches(0.7), y, Inches(2.0), box_h,
                 who, size=15, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # 内容
        add_round_rect(slide, Inches(2.85), y, Inches(9.85), box_h,
                       fill=NAVY_PALE if i % 2 == 0 else WHITE,
                       line=NAVY_LIGHT, line_w=1.0, radius=0.06)
        add_text(slide, Inches(3.05), y, Inches(9.5), box_h,
                 text, size=13, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        # 矢印
        if i < len(flows) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW,
                                           Inches(1.55), y + box_h - Inches(0.02),
                                           Inches(0.3), Inches(0.18))
            set_fill(arrow, NAVY_LIGHT); set_no_line(arrow)

    message_box(slide, Inches(6.85),
                "メッセージ：工房は『プロンプトを作る場所』。画像生成は新規チャット（or別Project）で行う",
                left=Inches(0.6), width=Inches(12.1))

slide_27()

# ============================================================
# スライド28: 演習②
# ============================================================
slide28_steps = [
    "作ったプロンプト工房Projectを開く",
    "自分が作りたい画像を抽象的に伝える（例：「次回セミナーの表紙、何かいい感じのを」）",
    "ヒアリングに答える",
    "出てきたプロンプトをコピー",
    "新規チャットでImages機能に貼り付けて画像生成",
]
slide28_checks = [
    "プロンプト工房Projectが作れた",
    "ヒアリングを経てプロンプトが返ってきた",
    "そのプロンプトで画像が生成できた",
    "生成画像が想定したテイストに近い",
]
exercise_slide(28, "演習②",
               "プロンプト工房を実際に使って画像を生成する",
               slide28_steps, slide28_checks, "7分",
               section="Part 2｜演習")

# ============================================================
# スライド29: セクション ― Part 3
# ============================================================
add_section_slide("PART 3", "同一テイストの複数枚\n＆ 生成後の編集", 29)

# ============================================================
# スライド30: 同一テイストで複数枚 ― 3つの方法
# ============================================================
def slide_30():
    slide = add_slide()
    add_header(slide, section_label="Part 3｜複数枚＆編集", page_num=30)
    add_title_block(slide, "同一テイストで複数枚を揃える ― 3つの方法",
                    "おすすめは『C → A』の組み合わせ")

    methods = [
        ("A", "同じ会話内で連続生成",
         "1枚目を作った同じチャットで\n「同じテイストで○○も作って」",
         "直前の生成結果を参考にしてくれる",
         NAVY_LIGHT),
        ("B", "1枚目を「参考画像」として添付",
         "新しいチャットで参考画像を添付して\n「同じスタイルで○○を作って」",
         "別の会話でも明示的に指定可",
         NAVY),
        ("C", "プロンプト工房に相談",
         "工房に\n「先ほどのテーマで挿絵も2枚追加で」",
         "共通要素を固定したバリエーションが返る",
         NAVY_DARK),
    ]
    y = Inches(1.95)
    box_h = Inches(4.4)
    box_w = Inches(4.0)
    gap = Inches(0.07)
    total = box_w * 3 + gap * 2
    start_x = (SLIDE_W - total) / 2
    for i, (label, title, how, feat, c) in enumerate(methods):
        x = start_x + (box_w + gap) * i
        add_round_rect(slide, x, y, box_w, box_h, fill=WHITE, line=c, line_w=2.5, radius=0.05)
        # ヘッダー
        add_rect(slide, x, y, box_w, Inches(0.85), fill=c)
        add_text(slide, x + Inches(0.2), y, Inches(0.7), Inches(0.85),
                 label, size=36, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + Inches(1.0), y, box_w - Inches(1.1), Inches(0.85),
                 title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        # やり方
        add_text(slide, x + Inches(0.2), y + Inches(1.0), box_w - Inches(0.4), Inches(0.4),
                 "▶ やり方", size=12, bold=True, color=NAVY_LIGHT)
        add_text(slide, x + Inches(0.2), y + Inches(1.4), box_w - Inches(0.4), Inches(1.2),
                 how, size=13, color=GRAY_TEXT)
        # 特徴
        add_text(slide, x + Inches(0.2), y + Inches(2.7), box_w - Inches(0.4), Inches(0.4),
                 "▶ 特徴", size=12, bold=True, color=NAVY_LIGHT)
        add_text(slide, x + Inches(0.2), y + Inches(3.1), box_w - Inches(0.4), Inches(1.2),
                 feat, size=13, color=NAVY_DARK)

    # 推奨組み合わせの強調
    add_round_rect(slide, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.75),
                   fill=NAVY, line=NAVY, line_w=2.0, radius=0.1)
    add_text(slide, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.75),
             "★  おすすめ：方法 C  →  出てきたプロンプトを 方法 A で使う  ＝  最も安定する",
             size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

slide_30()

# ============================================================
# スライド31: 生成後の編集 ― よく使う指示パターン
# ============================================================
def slide_31():
    slide = add_slide()
    add_header(slide, section_label="Part 3｜複数枚＆編集", page_num=31)
    add_title_block(slide, "生成後の編集 ― よく使う指示パターン",
                    "1回の指示で多くを変えず、小さく繰り返す")

    header = ["目的", "指示の例"]
    rows = [
        ("色を変える",      "「背景の色を青系に変えて」"),
        ("要素を追加",      "「左上にカップを追加して」"),
        ("要素を削除",      "「右側の人物を消して」"),
        ("スタイル変換",    "「同じ構図で水彩画風にして」"),
        ("文字を変える",    "「タイトル文字を“Spring Sale”に変えて」"),
    ]
    add_table(slide, Inches(1.0), Inches(1.95), Inches(11.3), Inches(3.3),
              header, rows, col_widths=[3.5, 7.8],
              header_h=Inches(0.5), row_h=Inches(0.55),
              header_size=15, body_size=14,
              align_cols=['L', 'L'])

    # 編集サイクル概念図
    add_text(slide, Inches(0.6), Inches(5.45), Inches(12.1), Inches(0.4),
             "▼  編集サイクル ― 小さく回す", size=13, bold=True, color=NAVY_DARK)
    cycle = ["生成", "1箇所だけ指示", "確認", "次の指示"]
    cw = Inches(2.6)
    aw = Inches(0.3)
    total = cw * 4 + aw * 3
    sx = (SLIDE_W - total) / 2
    yy = Inches(5.9)
    cyc_h = Inches(0.6)
    for i, t in enumerate(cycle):
        x = sx + (cw + aw) * i
        add_round_rect(slide, x, yy, cw, cyc_h,
                       fill=NAVY_PALE, line=NAVY, line_w=1.5, radius=0.2)
        add_text(slide, x, yy, cw, cyc_h,
                 t, size=14, bold=True, color=NAVY_DARK,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 3:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                           x + cw + Inches(0.01), yy + Inches(0.18),
                                           aw - Inches(0.02), Inches(0.25))
            set_fill(arrow, NAVY); set_no_line(arrow)

    message_box(slide, Inches(6.7),
                "メッセージ：1回の指示で多くを変えず、小さく繰り返すのがコツ",
                left=Inches(0.6), width=Inches(12.1))

slide_31()

# ============================================================
# スライド32: 精度を上げるコツ ― 変えない部分を明示
# ============================================================
def slide_32():
    slide = add_slide()
    add_header(slide, section_label="Part 3｜複数枚＆編集", page_num=32)
    add_title_block(slide, "⚠️ 精度を上げるコツ ― 「変えない部分」を明示する",
                    "「変える部分」より「変えない部分」を伝えるほうが効く")

    # 悪い例
    add_round_rect(slide, Inches(0.6), Inches(2.0), Inches(12.1), Inches(2.0),
                   fill=WHITE, line=NAVY_LIGHT, line_w=2.0, radius=0.05)
    add_rect(slide, Inches(0.6), Inches(2.0), Inches(2.5), Inches(2.0), fill=NAVY_LIGHT)
    add_text(slide, Inches(0.6), Inches(2.0), Inches(2.5), Inches(2.0),
             "❌\n悪い例", size=22, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(3.3), Inches(2.15), Inches(9.2), Inches(0.7),
             "「背景を変えて」", size=22, bold=True, color=NAVY_DARK)
    add_text(slide, Inches(3.3), Inches(2.85), Inches(9.2), Inches(1.0),
             "→ 全体が変わってしまうことがある（人物の表情・構図まで変わる）",
             size=14, color=GRAY_TEXT)

    # 良い例
    add_round_rect(slide, Inches(0.6), Inches(4.2), Inches(12.1), Inches(2.0),
                   fill=WHITE, line=NAVY, line_w=2.5, radius=0.05)
    add_rect(slide, Inches(0.6), Inches(4.2), Inches(2.5), Inches(2.0), fill=NAVY)
    add_text(slide, Inches(0.6), Inches(4.2), Inches(2.5), Inches(2.0),
             "✓\n良い例", size=22, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_runs(slide, Inches(3.3), Inches(4.32), Inches(9.2), Inches(0.7),
             [
                ("「背景の色だけを青系に変えて、", 18, True, NAVY_DARK),
                ("人物と構図はそのまま", 18, True, ACCENT_GOLD),
                ("にして」", 18, True, NAVY_DARK),
             ])
    add_text(slide, Inches(3.3), Inches(5.05), Inches(9.2), Inches(1.0),
             "→ 変えたいところだけ変わる（変えない部分を明示することで保持される）",
             size=14, color=NAVY_DARK)

    message_box(slide, Inches(6.45),
                "メッセージ：「変える部分」より「変えない部分」を伝えるほうが効く",
                left=Inches(0.6), width=Inches(12.1))

slide_32()

# ============================================================
# スライド33: 演習③
# ============================================================
def slide_33():
    slide = add_slide()
    add_header(slide, section_label="Part 3｜演習", page_num=33)
    # 演習バッジ
    add_rect(slide, Inches(0.5), Inches(0.65), Inches(0.18), Inches(0.85), fill=ACCENT_GOLD)
    add_runs(slide, Inches(0.78), Inches(0.55), Inches(11.5), Inches(0.7),
             [
                ("演習③", 30, True, NAVY_DARK),
                ("   ", 30, False, NAVY_DARK),
                ("⏱ 5分", 18, True, ACCENT_GOLD),
             ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.78), Inches(1.20), Inches(11.5), Inches(0.4),
             "演習②で作った画像を発展させる（A・Bどちらか選択）",
             size=14, bold=False, color=NAVY_LIGHT)
    add_rect(slide, Inches(0.5), Inches(1.62), Inches(12.3), Inches(0.02), fill=NAVY_LIGHT)

    # A・B 選択
    box_w = Inches(5.95)
    box_h = Inches(3.5)
    y = Inches(2.0)
    # A
    add_round_rect(slide, Inches(0.6), y, box_w, box_h, fill=WHITE, line=NAVY, line_w=2.5, radius=0.05)
    add_rect(slide, Inches(0.6), y, box_w, Inches(0.7), fill=NAVY)
    add_text(slide, Inches(0.6), y, box_w, Inches(0.7),
             "A：バリエーション追加", size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.85), y + Inches(0.95), box_w - Inches(0.5), Inches(2.0),
             "演習②で作った画像と同じテイストで、別バージョンを1枚生成", size=15, color=GRAY_TEXT)
    add_text(slide, Inches(0.85), y + Inches(2.0), box_w - Inches(0.5), Inches(1.4),
             "例：背景を別の季節に／登場キャラを別バージョンに／別シーンに",
             size=12, color=NAVY_LIGHT)

    # B
    add_round_rect(slide, Inches(6.75), y, box_w, box_h, fill=WHITE, line=NAVY, line_w=2.5, radius=0.05)
    add_rect(slide, Inches(6.75), y, box_w, Inches(0.7), fill=NAVY)
    add_text(slide, Inches(6.75), y, box_w, Inches(0.7),
             "B：部分修正", size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(7.0), y + Inches(0.95), box_w - Inches(0.5), Inches(0.6),
             "演習②で作った画像を、編集指示で修正", size=15, color=GRAY_TEXT)
    add_text(slide, Inches(7.0), y + Inches(1.7), box_w - Inches(0.5), Inches(1.7),
             "例：「背景の色を変えて」「要素を1つ追加」「別スタイルに変換」",
             size=12, color=NAVY_LIGHT)

    # チェックポイント
    add_text(slide, Inches(0.6), Inches(5.7), Inches(12.1), Inches(0.4),
             "▼  チェックポイント", size=14, bold=True, color=NAVY_DARK)
    add_round_rect(slide, Inches(0.6), Inches(6.15), Inches(12.1), Inches(1.05),
                   fill=NAVY_PALE, line=NAVY, line_w=1.5, radius=0.05)
    checks = [
        "バリエーション追加 or 部分修正ができた",
        "「変えない部分」が崩れずに残っている",
    ]
    for i, c in enumerate(checks):
        x = Inches(1.0) + (Inches(5.85) * i)
        cb = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(6.5), Inches(0.32), Inches(0.32))
        set_fill(cb, WHITE); set_line(cb, NAVY, 1.5)
        add_text(slide, x + Inches(0.45), Inches(6.4), Inches(5.5), Inches(0.55),
                 c, size=14, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)

slide_33()

# ============================================================
# スライド34: セクション ― Part 4
# ============================================================
add_section_slide("PART 4", "Claudeに埋め込む\n― Vol.14連携", 34)

# ============================================================
# スライド35: なぜBase64で埋め込むのか
# ============================================================
def slide_35():
    slide = add_slide()
    add_header(slide, section_label="Part 4｜Claude連携", page_num=35)
    add_title_block(slide, "なぜBase64で埋め込むのか",
                    "1ファイルで完結 ― 共有・印刷・PDF化がシンプル")

    # 比較概念図：通常 vs Base64埋め込み
    y = Inches(2.0)
    col_h = Inches(2.7)
    col_w = Inches(5.95)
    # 通常
    add_round_rect(slide, Inches(0.6), y, col_w, col_h, fill=WHITE, line=NAVY_LIGHT, line_w=1.8, radius=0.05)
    add_rect(slide, Inches(0.6), y, col_w, Inches(0.55), fill=NAVY_LIGHT)
    add_text(slide, Inches(0.6), y, col_w, Inches(0.55),
             "通常のHTML", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    files_n = ["📄 report.html", "🖼 image1.png", "🖼 image2.png"]
    for i, f in enumerate(files_n):
        x = Inches(0.85) + Inches(1.85) * i
        add_round_rect(slide, x, y + Inches(0.85), Inches(1.6), Inches(1.2),
                       fill=NAVY_PALE, line=NAVY_LIGHT, line_w=1.0, radius=0.08)
        add_text(slide, x, y + Inches(0.85), Inches(1.6), Inches(1.2),
                 f, size=12, color=NAVY_DARK,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.6), y + Inches(2.15), col_w, Inches(0.5),
             "→  3ファイルを管理する必要がある", size=13, bold=True, color=NAVY_LIGHT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Base64
    add_round_rect(slide, Inches(6.75), y, col_w, col_h, fill=WHITE, line=NAVY, line_w=2.5, radius=0.05)
    add_rect(slide, Inches(6.75), y, col_w, Inches(0.55), fill=NAVY)
    add_text(slide, Inches(6.75), y, col_w, Inches(0.55),
             "Base64埋め込みHTML", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # 1つの大きいファイルアイコン（中に画像が入っているイメージ）
    add_round_rect(slide, Inches(7.65), y + Inches(0.85), Inches(4.15), Inches(1.2),
                   fill=NAVY_PALE, line=NAVY, line_w=2.0, radius=0.08)
    add_text(slide, Inches(7.65), y + Inches(0.85), Inches(4.15), Inches(0.5),
             "📄  report.html", size=14, bold=True, color=NAVY_DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(7.65), y + Inches(1.3), Inches(4.15), Inches(0.7),
             "（画像データもこの中に Base64で内包）", size=11, color=GRAY_TEXT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(6.75), y + Inches(2.15), col_w, Inches(0.5),
             "→  1ファイルで完結", size=13, bold=True, color=NAVY,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # メリットリスト
    add_text(slide, Inches(0.6), Inches(5.0), Inches(12.1), Inches(0.4),
             "▼  メリット", size=14, bold=True, color=NAVY_DARK)
    benefits = [
        "HTML 1ファイルで完結 → 共有・印刷・PDF化がシンプル",
        "画像のリンク切れがない（他環境でも崩れない）",
        "メールでも1ファイル送るだけ",
    ]
    for i, b in enumerate(benefits):
        yy = Inches(5.45) + Inches(0.6) * i
        add_rect(slide, Inches(0.6), yy + Inches(0.15), Inches(0.3), Inches(0.3), fill=NAVY)
        add_text(slide, Inches(0.6), yy + Inches(0.15), Inches(0.3), Inches(0.3),
                 "✓", size=14, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, Inches(1.0), yy, Inches(11.5), Inches(0.55),
                 b, size=14, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)

slide_35()

# ============================================================
# スライド36: Claudeに画像を渡す → Base64で埋め込ませる
# ============================================================
def slide_36():
    slide = add_slide()
    add_header(slide, section_label="Part 4｜Claude連携", page_num=36)
    add_title_block(slide, "Claudeに画像を渡す → Base64で埋め込ませる",
                    "5ステップで成果物が完成")

    steps = [
        ("1", "ファイルアプリ（または写真アプリ）から、埋め込みたい画像を準備"),
        ("2", "Claudeを開いて新規チャット"),
        ("3", "添付アイコンから画像を添付"),
        ("4", "指示文を入力（下のテキストをコピペ可）"),
        ("5", "ClaudeがHTMLを生成 → ダウンロード → iPad表示確認 →（必要ならPDF変換）"),
    ]
    y = Inches(1.95)
    sh = Inches(0.55)
    gap = Inches(0.08)
    for i, (n, t) in enumerate(steps):
        yi = y + (sh + gap) * i
        add_rect(slide, Inches(0.6), yi, Inches(0.7), sh, fill=NAVY)
        add_text(slide, Inches(0.6), yi, Inches(0.7), sh,
                 n, size=18, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_rect(slide, Inches(1.35), yi, Inches(11.35), sh, fill=NAVY_PALE)
        add_text(slide, Inches(1.5), yi, Inches(11.2), sh,
                 t, size=13, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)

    # 指示文ボックス（コードっぽく）
    code_y = Inches(5.25)
    add_rect(slide, Inches(0.6), code_y, Inches(12.1), Inches(0.4), fill=NAVY_DARK)
    add_text(slide, Inches(0.8), code_y, Inches(12.1), Inches(0.4),
             "📝  指示文（コピペ用）", size=12, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, Inches(0.6), code_y + Inches(0.4), Inches(12.1), Inches(1.7),
             fill=NAVY_PALE, line=NAVY_DARK, line_w=1.5)
    code_text = (
        "この画像を使って、簡単な自己紹介HTMLを作ってください。\n"
        "画像はBase64でHTMLに埋め込んで、表紙画像として配置してください。\n"
        "タイトルは「自己紹介」、本文は短くて構いません。"
    )
    add_text(slide, Inches(0.85), code_y + Inches(0.5), Inches(11.6), Inches(1.5),
             code_text, size=14, color=NAVY_DARK, font='Menlo')

slide_36()

# ============================================================
# スライド37: 演習④
# ============================================================
slide37_steps = [
    "Part 1〜3で作った画像のうち、お気に入りを1枚選ぶ",
    "Claudeに短いHTMLドキュメントを作らせる（例：本講義の感想 / 自己紹介 / プロフィール紹介）",
    "上の画像をBase64で埋め込ませる",
    "生成されたHTMLをiPadで表示確認",
    "（時間があれば）PDFに変換",
]
slide37_checks = [
    "HTMLに画像がBase64で埋め込まれている",
    "iPadで開いたときに画像が表示される",
    "PDF化したときも画像が表示される",
]
exercise_slide(37, "演習④（統合演習）",
               "今日作った画像を、Claudeのドキュメントに組み込んで「成果物」として仕上げる",
               slide37_steps, slide37_checks, "10分",
               section="Part 4｜統合演習")

# ============================================================
# スライド38: 今日学んだこと
# ============================================================
def slide_38():
    slide = add_slide()
    add_header(slide, section_label="まとめ", page_num=38)
    add_title_block(slide, "今日学んだこと（ゴール8項目の振り返り）",
                    "全部できるようになった ＝ ゴール達成")

    items = [
        ("1", "ChatGPTのImages機能で画像を生成できる",         "Part 1"),
        ("2", "用途別にアスペクト比を使い分けられる",            "Part 1"),
        ("3", "プロンプト工房でプロンプトを生成できる",          "Part 2"),
        ("4", "同一テイストで複数枚を揃えられる",                "Part 3"),
        ("5", "チャットで生成後編集ができる",                    "Part 3"),
        ("6", "写真App／ファイルAppに保存できる",               "Part 1"),
        ("7", "ClaudeのHTMLにBase64で埋め込める",              "Part 4"),
        ("8", "苦手領域を理解し、被写体抜き出しで補完できる",     "Part 0／1"),
    ]
    # 2列レイアウト
    col_w = Inches(5.95)
    row_h = Inches(0.62)
    y0 = Inches(2.0)
    for i, (n, t, src) in enumerate(items):
        col = i // 4
        row = i % 4
        x = Inches(0.6) + (col_w + Inches(0.2)) * col
        y = y0 + (row_h + Inches(0.1)) * row
        add_round_rect(slide, x, y, col_w, row_h,
                       fill=NAVY_PALE, line=NAVY_LIGHT, line_w=1.0, radius=0.1)
        # 番号
        add_rect(slide, x + Inches(0.05), y + Inches(0.1), Inches(0.45), Inches(0.42), fill=NAVY)
        add_text(slide, x + Inches(0.05), y + Inches(0.1), Inches(0.45), Inches(0.42),
                 n, size=15, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # 本文
        add_text(slide, x + Inches(0.6), y, col_w - Inches(2.0), row_h,
                 t, size=13, color=NAVY_DARK, anchor=MSO_ANCHOR.MIDDLE)
        # 出典
        add_rect(slide, x + col_w - Inches(1.3), y + Inches(0.13), Inches(1.2), Inches(0.36),
                 fill=NAVY)
        add_text(slide, x + col_w - Inches(1.3), y + Inches(0.13), Inches(1.2), Inches(0.36),
                 src, size=11, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # 達成バナー
    add_rect(slide, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.85), fill=NAVY)
    add_text(slide, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.85),
             "🎉  おめでとうございます！  Vol.16のゴール8項目をすべて達成 🎉",
             size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

slide_38()

# ============================================================
# スライド39: 使い道早見表
# ============================================================
def slide_39():
    slide = add_slide()
    add_header(slide, section_label="まとめ", page_num=39)
    add_title_block(slide, "作った画像はこんな場面で使える ― 使い道早見表",
                    "迷ったら KeynoteかGoodNotes に貼る")

    header = ["ジャンル", "主な使い道", "配置先"]
    rows = [
        ("提案書／企画書の表紙",      "会議・商談の第一印象を上げる", "Keynote / PowerPoint"),
        ("業務マニュアル用アイコン",   "マニュアル・FAQの統一感",    "Word / Pages / Notion"),
        ("誕生日招待状／季節カード",   "親戚・友人・取引先への挨拶",  "LINE / メール / 印刷"),
        ("町内会・PTAポスター",       "地域・学校への告知",          "掲示板 / LINEグループ"),
        ("ブログアイキャッチ",        "記事の第一印象アップ",        "ブログ / note"),
        ("Instagram投稿／ヘッダー",  "SNS発信・お店紹介",          "Instagram / X"),
        ("子どもの学習プリント",      "お子さんの勉強サポート",      "印刷 / GoodNotes"),
        ("資格学習・社内研修の素材",  "大人の学び・OJT資料",         "GoodNotes / Keynote"),
    ]
    add_table(slide, Inches(0.5), Inches(1.95), Inches(12.3), Inches(4.5),
              header, rows, col_widths=[3.7, 4.6, 4.0],
              header_h=Inches(0.4), row_h=Inches(0.5),
              header_size=13, body_size=12,
              align_cols=['L', 'L', 'L'])

    message_box(slide, Inches(6.55),
                "メッセージ：迷ったら【KeynoteかGoodNotesに貼る】が一番つぶしが効く（PDF化で共有・印刷・メール添付すべて対応）",
                left=Inches(0.5), width=Inches(12.3))

slide_39()

# ============================================================
# スライド40: エンドスライド
# ============================================================
def slide_40():
    slide = add_slide()
    # 全面紺色
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY)
    # 装飾線
    add_rect(slide, Inches(2.5), Inches(2.4), Inches(8.3), Inches(0.06), fill=NAVY_PALE)
    # メイン
    add_text(slide, 0, Inches(2.7), SLIDE_W, Inches(1.7),
             "ありがとうございました！", size=72, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # サブ
    add_text(slide, 0, Inches(4.5), SLIDE_W, Inches(0.6),
             "質問・感想はチャットでお気軽にどうぞ", size=22, color=NAVY_PALE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # 装飾線
    add_rect(slide, Inches(4.5), Inches(5.3), Inches(4.3), Inches(0.04), fill=NAVY_LIGHT)
    # フッター
    add_text(slide, 0, Inches(5.7), SLIDE_W, Inches(0.4),
             "Vol.16  ｜  AIデザイン作成  ―  GPT Image 2でデザインを自動化する",
             size=14, color=NAVY_PALE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0, Inches(6.2), SLIDE_W, Inches(0.4),
             "AI活用セミナーシリーズ  ―  Phase 4：AIによる創造とアウトプット",
             size=12, color=NAVY_LIGHT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0, Inches(7.05), SLIDE_W, Inches(0.4),
             "40 / 40", size=10, color=NAVY_PALE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

slide_40()

# ============================================================
# 保存
# ============================================================
out = "/home/user/test/Vol16_AIデザイン作成.pptx"
prs.save(out)
print(f"Saved: {out}")
print(f"Total slides: {len(prs.slides)}")
